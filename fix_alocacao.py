import openpyxl

ARQ = "Simulador_FII.xlsx"

# Garante que a coluna CHAVE da aba Perfis esteja preenchida,
# evitando o erro #N/D no PROCV da alocação por tipo de FII.
wb = openpyxl.load_workbook(ARQ)
perfis = wb["Perfis"]

for r in range(3, 21):
    perfis[f"A{r}"] = f"=B{r}&\"-\"&C{r}"

# Ajusta o texto de dica para referenciar a célula correta do perfil
ws = wb["Simulador"]
if ws["K52"].value and "H21" in str(ws["K52"].value):
    ws["K52"] = ("Dica: altere os valores em azul para simular cenários e use o menu "
                 "suspenso em 'Perfil de risco' (L41) para recalcular a alocação por tipo de FII.")

wb.save(ARQ)
print("CHAVE da aba Perfis garantida e texto de dica ajustado.")
