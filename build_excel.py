import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

# ----------------------------------------------------------------------------
# Paleta de cores (tema "Fundos Imobiliários")
# ----------------------------------------------------------------------------
AZUL_ESCURO = "1F3864"
AZUL_MEDIO = "2E5496"
AZUL_CLARO = "D9E1F2"
VERDE = "375623"
VERDE_CLARO = "E2EFDA"
CINZA = "F2F2F2"
BRANCO = "FFFFFF"
DOURADO = "BF9000"

REAL = 'R$ #,##0.00'
PCT = '0.00%'
PCT1 = '0.0%'

wb = openpyxl.Workbook()

# ============================================================================
# ESTILOS REUTILIZÁVEIS
# ============================================================================
titulo = Font(name="Calibri", size=18, bold=True, color=BRANCO)
subtitulo = Font(name="Calibri", size=13, bold=True, color=BRANCO)
secao = Font(name="Calibri", size=12, bold=True, color=BRANCO)
label = Font(name="Calibri", size=11, bold=False, color="000000")
label_bold = Font(name="Calibri", size=11, bold=True, color="000000")
resultado = Font(name="Calibri", size=11, bold=True, color=VERDE)
valor_in = Font(name="Calibri", size=11, bold=True, color=AZUL_ESCURO)

preench_azul = PatternFill("solid", fgColor=AZUL_ESCURO)
preench_azul_med = PatternFill("solid", fgColor=AZUL_MEDIO)
preench_verde = PatternFill("solid", fgColor=VERDE)
preench_claro = PatternFill("solid", fgColor=AZUL_CLARO)
preench_verde_claro = PatternFill("solid", fgColor=VERDE_CLARO)
preench_cinza = PatternFill("solid", fgColor=CINZA)

borda = Side(style="thin", color="BFBFBF")
box = Border(left=borda, right=borda, top=borda, bottom=borda)

centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
esq = Alignment(horizontal="left", vertical="center")
dir_ = Alignment(horizontal="right", vertical="center")

# ============================================================================
# PLANILHA 1 - SIMULADOR
# ============================================================================
ws = wb.active
ws.title = "Simulador"
ws.sheet_view.showGridLines = False

# Larguras
ws.column_dimensions["A"].width = 4
ws.column_dimensions["B"].width = 40
ws.column_dimensions["C"].width = 22
ws.column_dimensions["D"].width = 22
ws.column_dimensions["E"].width = 18

# ---------- Cabeçalho ----------
ws.merge_cells("B2:E2")
c = ws["B2"]; c.value = "SIMULADOR DE INVESTIMENTOS EM FUNDOS IMOBILIÁRIOS (FII)"
c.font = titulo; c.fill = preench_azul; c.alignment = centro
ws.row_dimensions[2].height = 34

ws.merge_cells("B3:E3")
c = ws["B3"]; c.value = "Ferramenta de apoio à decisão: patrimônio acumulado, dividendos mensais e alocação por perfil"
c.font = Font(size=10, italic=True, color="404040"); c.fill = preench_cinza; c.alignment = centro
ws.row_dimensions[3].height = 20

# ---------- CONFIGURAÇÕES ----------
ws.merge_cells("B5:E5")
c = ws["B5"]; c.value = "CONFIGURAÇÕES DO INVESTIDOR"; c.font = secao; c.fill = preench_azul_med; c.alignment = esq

ws["B6"] = "Salário"; ws["B6"].font = label
ws["D6"] = 2000; ws["D6"].number_format = REAL; ws["D6"].font = valor_in; ws["D6"].fill = preench_claro; ws["D6"].border = box
ws["D6"].alignment = dir_

ws["B7"] = "Rendimento médio da carteira (dividendos)"; ws["B7"].font = label
ws["D7"] = 0.006; ws["D7"].number_format = PCT; ws["D7"].font = valor_in; ws["D7"].fill = preench_claro; ws["D7"].border = box
ws["D7"].alignment = dir_

ws["B8"] = "Sugestão de investimento (30% do salário)"; ws["B8"].font = label
ws["D8"] = "=D6*30%"; ws["D8"].number_format = REAL; ws["D8"].font = resultado; ws["D8"].fill = preench_verde_claro; ws["D8"].border = box
ws["D8"].alignment = dir_

# ---------- INVESTIMENTO MENSAL ----------
ws.merge_cells("B10:E10")
c = ws["B10"]; c.value = "PARÂMETROS DA SIMULAÇÃO"; c.font = secao; c.fill = preench_azul_med; c.alignment = esq

ws["B11"] = "Quanto investir por mês?"; ws["B11"].font = label_bold
ws["D11"] = 200; ws["D11"].number_format = REAL; ws["D11"].font = valor_in; ws["D11"].fill = preench_claro; ws["D11"].border = box
ws["D11"].alignment = dir_

ws["B12"] = "Por quantos anos?"; ws["B12"].font = label_bold
ws["D12"] = 5; ws["D12"].number_format = '0'; ws["D12"].font = valor_in; ws["D12"].fill = preench_claro; ws["D12"].border = box
ws["D12"].alignment = dir_

ws["B13"] = "Taxa de rendimento mensal (valorização)?"; ws["B13"].font = label_bold
ws["D13"] = 0.01079; ws["D13"].number_format = PCT; ws["D13"].font = valor_in; ws["D13"].fill = preench_claro; ws["D13"].border = box
ws["D13"].alignment = dir_

ws["B14"] = "Patrimônio acumulado"; ws["B14"].font = label_bold
ws["D14"] = "=FV(taxa_mensal,qtd_anos*12,aporte*-1)"; ws["D14"].number_format = REAL
ws["D14"].font = resultado; ws["D14"].fill = preench_verde_claro; ws["D14"].border = box; ws["D14"].alignment = dir_

ws["B15"] = "Dividendos mensais estimados"; ws["B15"].font = label_bold
ws["D15"] = "=patrimonio*rendimento_carteira"; ws["D15"].number_format = REAL
ws["D15"].font = resultado; ws["D15"].fill = preench_verde_claro; ws["D15"].border = box; ws["D15"].alignment = dir_

ws["B16"] = "Total efetivamente investido"; ws["B16"].font = label
ws["D16"] = "=aporte*qtd_anos*12"; ws["D16"].number_format = REAL
ws["D16"].font = label; ws["D16"].fill = preench_cinza; ws["D16"].border = box; ws["D16"].alignment = dir_

ws["B17"] = "Rendimento (juros) acumulado"; ws["B17"].font = label
ws["D17"] = "=patrimonio-D16"; ws["D17"].number_format = REAL
ws["D17"].font = label; ws["D17"].fill = preench_cinza; ws["D17"].border = box; ws["D17"].alignment = dir_

# ---------- CENÁRIOS ----------
ws.merge_cells("B19:E19")
c = ws["B19"]; c.value = "CENÁRIOS - EVOLUÇÃO DO PATRIMÔNIO"; c.font = secao; c.fill = preench_azul_med; c.alignment = esq

ws["A20"] = "Ano"; ws["B20"] = "Descrição"
ws["C20"] = "Patrimônio"; ws["D20"] = "Dividendo Mensal"
for col in ["A", "B", "C", "D"]:
    cc = ws[f"{col}20"]; cc.font = secao; cc.fill = preench_azul_med; cc.alignment = centro; cc.border = box

anos = [2, 5, 10, 20, 30]
r = 21
for i, a in enumerate(anos):
    ws[f"A{r}"] = a; ws[f"A{r}"].alignment = centro; ws[f"A{r}"].border = box
    ws[f"B{r}"] = f"Quanto em {a} anos?"; ws[f"B{r}"].font = label; ws[f"B{r}"].border = box
    ws[f"C{r}"] = f"=FV($D$13,A{r}*12,$D$11*-1)"; ws[f"C{r}"].number_format = REAL
    ws[f"C{r}"].font = resultado; ws[f"C{r}"].border = box; ws[f"C{r}"].alignment = dir_
    ws[f"D{r}"] = f"=C{r}*rendimento_carteira"; ws[f"D{r}"].number_format = REAL
    ws[f"D{r}"].font = label; ws[f"D{r}"].border = box; ws[f"D{r}"].alignment = dir_
    if i % 2 == 0:
        for col in ["A", "B", "C", "D"]:
            if ws[f"{col}{r}"].fill.fgColor.rgb in (None, "00000000"):
                ws[f"{col}{r}"].fill = preench_cinza
    r += 1

# ---------- PROJEÇÃO MENSAL (tabela para gráfico) ----------
ws.merge_cells("B28:E28")
c = ws["B28"]; c.value = "PROJEÇÃO MENSAL (para visualização gráfica)"; c.font = secao; c.fill = preench_azul_med; c.alignment = esq

ws["B29"] = "Mês"; ws["C29"] = "Patrimônio (R$)"
ws["B29"].font = label_bold; ws["C29"].font = label_bold
ws["B29"].fill = preench_claro; ws["C29"].fill = preench_claro
ws["B29"].border = box; ws["C29"].border = box
ws["B29"].alignment = centro; ws["C29"].alignment = centro

total_meses = "qtd_anos*12"
for m in range(1, 61):  # até 60 meses (5 anos) para o gráfico padrão
    rr = 29 + m
    ws[f"B{rr}"] = m
    ws[f"C{rr}"] = f"=FV($D$13,B{rr},$D$11*-1)"
    ws[f"B{rr}"].alignment = centro; ws[f"B{rr}"].border = box
    ws[f"C{rr}"].number_format = REAL; ws[f"C{rr}"].border = box; ws[f"C{rr}"].alignment = dir_

# ---------- PERFIL DE ALOCAÇÃO ----------
ws.merge_cells("B91:E91")
c = ws["B91"]; c.value = "PERFIL DE ALOCAÇÃO DA CARTEIRA"; c.font = secao; c.fill = preench_azul_med; c.alignment = esq

ws["B92"] = "Perfil de risco"; ws["B92"].font = label_bold
ws["C92"] = "Moderado"; ws["C92"].font = valor_in; ws["C92"].fill = preench_claro; ws["C92"].border = box; ws["C92"].alignment = centro

ws["B93"] = "Valor a ser investido por mês"; ws["B93"].font = label
ws["C93"] = "=aporte"; ws["C93"].number_format = REAL; ws["C93"].font = label; ws["C93"].border = box; ws["C93"].alignment = dir_

ws["B95"] = "Tipo de FII"; ws["C95"] = "Percentual Sugerido"; ws["D95"] = "Valor Mensal (R$)"
for col in ["B", "C", "D"]:
    cc = ws[f"{col}95"]; cc.font = secao; cc.fill = preench_azul_med; cc.alignment = centro; cc.border = box

tipos = ["PAPEL", "TIJOLO", "HÍBRIDOS", "FOFs", "DESENVOLVIMENTO", "HOTELARIAS"]
r = 96
for t in tipos:
    ws[f"B{r}"] = t; ws[f"B{r}"].font = label; ws[f"B{r}"].border = box
    ws[f"C{r}"] = f'=VLOOKUP($C$92&"-"&B{r},Perfis!$A:$D,4,FALSE)'
    ws[f"C{r}"].number_format = PCT1; ws[f"C{r}"].font = label; ws[f"C{r}"].border = box; ws[f"C{r}"].alignment = dir_
    ws[f"D{r}"] = f"=C{r}*$C$93"; ws[f"D{r}"].number_format = REAL; ws[f"D{r}"].font = label; ws[f"D{r}"].border = box; ws[f"D{r}"].alignment = dir_
    r += 1

# Rodapé
ws.merge_cells("B103:E103")
ws["B103"] = "Dica: altere os valores em azul para simular diferentes cenários. Os campos verdes são calculados automaticamente."
ws["B103"].font = Font(size=9, italic=True, color="808080")

# ----------------------------------------------------------------------------
# Gráfico de projeção (linha) - patrimônio nos primeiros 60 meses
# ----------------------------------------------------------------------------
chart = LineChart()
chart.title = "Evolução do Patrimônio (primeiros 60 meses)"
chart.y_axis.title = "Patrimônio (R$)"
chart.x_axis.title = "Mês"
chart.height = 8
chart.width = 16
data = Reference(ws, min_col=3, min_row=29, max_row=89)
cats = Reference(ws, min_col=2, min_row=30, max_row=89)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)
chart.legend = None
ws.add_chart(chart, "G2")

# ============================================================================
# PLANILHA 2 - PERFIS (tabela de referência)
# ============================================================================
ws2 = wb.create_sheet("Perfis")
ws2.sheet_view.showGridLines = False
for col, w in zip("ABCD", [30, 18, 18, 12]):
    ws2.column_dimensions[col].width = w

ws2.merge_cells("A1:D1")
c = ws2["A1"]; c.value = "TABELA DE PERFIS E ALOCAÇÃO POR TIPO DE FII"
c.font = secao; c.fill = preench_azul; c.alignment = centro
ws2.row_dimensions[1].height = 26

headers = ["CHAVE", "PERFIL", "TIPO DE FII", "%"]
for j, h in enumerate(headers):
    cc = ws2[f"{get_column_letter(j+1)}2"]; cc.value = h
    cc.font = secao; cc.fill = preench_azul_med; cc.alignment = centro; cc.border = box

perfis = [
    ("Conservador", [("PAPEL", 0.30), ("TIJOLO", 0.50), ("HÍBRIDOS", 0.10), ("FOFs", 0.10), ("DESENVOLVIMENTO", 0.0), ("HOTELARIAS", 0.0)]),
    ("Moderado",    [("PAPEL", 0.32), ("TIJOLO", 0.35), ("HÍBRIDOS", 0.08), ("FOFs", 0.05), ("DESENVOLVIMENTO", 0.10), ("HOTELARIAS", 0.10)]),
    ("Agressivo",   [("PAPEL", 0.50), ("TIJOLO", 0.10), ("HÍBRIDOS", 0.05), ("FOFs", 0.05), ("DESENVOLVIMENTO", 0.20), ("HOTELARIAS", 0.10)]),
]
r = 3
for perfil, itens in perfis:
    for tipo, pct in itens:
        ws2[f"A{r}"] = f"={perfil}&'-'&{tipo}" if False else f"={get_column_letter(2)}{r}&'-'&{get_column_letter(3)}{r}"
        ws2[f"B{r}"] = perfil
        ws2[f"C{r}"] = tipo
        ws2[f"D{r}"] = pct; ws2[f"D{r}"].number_format = PCT1
        for col in "ABCD":
            cc = ws2[f"{col}{r}"]; cc.border = box
            if col in "AB":
                cc.font = label
            else:
                cc.font = label
            if r % 2 == 1:
                cc.fill = preench_cinza
            cc.alignment = centro if col != "C" else esq
        r += 1

# Observação
ws2.merge_cells(f"A{r+1}:D{r+1}")
ws2[f"A{r+1}"] = "Os percentuais representam a sugestão de alocação do aporte mensal conforme o perfil de risco selecionado."
ws2[f"A{r+1}"].font = Font(size=9, italic=True, color="808080")

# ============================================================================
# NAMED RANGES utilitários da simulação
# ============================================================================
def add_name(name, ref):
    wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName(name, attr_text=ref))

add_name("salario", "Simulador!$D$6")
add_name("rendimento_carteira", "Simulador!$D$7")
add_name("sugestao_investimento", "Simulador!$D$8")
add_name("aporte", "Simulador!$D$11")
add_name("qtd_anos", "Simulador!$D$12")
add_name("taxa_mensal", "Simulador!$D$13")
add_name("patrimonio", "Simulador!$D$14")

# ============================================================================
wb.save("Simulador_FII.xlsx")
print("Arquivo 'Simulador_FII.xlsx' criado com sucesso.")
