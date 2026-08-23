import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice
from openpyxl.drawing.text import CharacterProperties

ARQ = "Simulador_FII.xlsx"
wb = openpyxl.load_workbook(ARQ)
ws = wb["Simulador"]
Perfis = wb["Perfis"]

REAL = 'R$ #,##0.00'

# ---------------------------------------------------------------------------
# 1) Dropdown de PERFIL DE RISCO em H21 (muda o perfil -> recalcula FII)
# ---------------------------------------------------------------------------
dv = DataValidation(type="list",
                    formula1='"Conservador,Moderado,Agressivo"',
                    allow_blank=False,
                    showDropDown=False)
dv.error = "Escolha um perfil: Conservador, Moderado ou Agressivo."
dv.errorTitle = "Perfil inválido"
dv.prompt = "Selecione o perfil de risco da carteira."
dv.promptTitle = "Perfil de Risco"
ws.add_data_validation(dv)
dv.add(ws["H21"])

# Destaca a célula de seleção de perfil
azul = PatternFill("solid", fgColor="D9E1F2")
ws["H21"].fill = azul
ws["H21"].font = Font(bold=True, color="1F3864", size=11)
ws["H21"].alignment = Alignment(horizontal="center", vertical="center")

# ---------------------------------------------------------------------------
# 2) Enriquecer a PROJEÇÃO MENSAL com Total Investido e Rendimento
#    (B=Mês, C=Patrimônio, D=Total Investido, E=Rendimento)
# ---------------------------------------------------------------------------
ws["D27"] = "Total Investido (R$)"
ws["E27"] = "Rendimento (R$)"
for col in ("D", "E"):
    cc = ws[f"{col}27"]
    cc.font = Font(bold=True, color="FFFFFF")
    cc.fill = PatternFill("solid", fgColor="2E5496")
    cc.alignment = Alignment(horizontal="center", vertical="center")
    cc.border = Border(bottom=Side(style="thin", color="BFBFBF"))

for r in range(28, 88):
    d = ws[f"D{r}"]
    d.value = f"=aporte*B{r}"
    d.number_format = REAL
    d.border = Border(left=Side(style="thin", color="D9D9D9"),
                      right=Side(style="thin", color="D9D9D9"))
    e = ws[f"E{r}"]
    e.value = f"=C{r}-D{r}"
    e.number_format = REAL
    e.border = Border(left=Side(style="thin", color="D9D9D9"),
                      right=Side(style="thin", color="D9D9D9"))

# ---------------------------------------------------------------------------
# 3) Remover o gráfico antigo e criar um modelo melhorado
# ---------------------------------------------------------------------------
ws._charts = []

chart = LineChart()
chart.title = "Evolução do Patrimônio, Investimento e Rendimento"
chart.style = 2
chart.height = 9
chart.width = 20

# Séries: Patrimônio (C), Total Investido (D), Rendimento (E)
data = Reference(ws, min_col=3, max_col=5, min_row=27, max_row=87)
cats = Reference(ws, min_col=2, min_row=28, max_row=87)
chart.add_data(data, titles_from_data=True)
chart.set_categories(cats)

# Eixos e títulos
chart.x_axis.title = "Mês"
chart.y_axis.title = "Valor (R$)"
chart.x_axis.delete = False
chart.y_axis.delete = False
chart.x_axis.scaling.orientation = "minMax"
chart.y_axis.numFmt = 'R$ #,##0'
chart.y_axis.majorGridlines = None

# Legenda embaixo
chart.legend.position = "b"

# Estilo por série: cor, espessura e marcadores
cores = ["1F3864", "A6A6A6", "375623"]  # azul, cinza, verde
for i, ser in enumerate(chart.series):
    ser.graphicalProperties.line = LineProperties(
        w=22000, solidFill=cores[i % 3])
    ser.graphicalProperties.line.dashStyle = "solid"
    # marcadores pequenos
    ser.marker.symbol = "circle"
    ser.marker.size = 3
    ser.smooth = False

# Fonte do título em negrito (quando suportado)
try:
    chart.title.tx.rich.p[0].r[0].rPr = CharacterProperties(b=True, sz=1400)
except Exception:
    pass

# Posicionar em área livre (coluna K)
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
anchor = OneCellAnchor(
    _from=AnchorMarker(col=10, colOff=0, row=1, rowOff=0),
    ext=XDRPositiveSize2D(cx=914400 * 18, cy=914400 * 9),
)
chart.anchor = anchor
ws.add_chart(chart)

# ---------------------------------------------------------------------------
# 4) Atualizar dica para reforçar o uso do perfil
# ---------------------------------------------------------------------------
ws["G32"] = ("Dica: altere os valores em azul para simular cenários e use o menu " 
             "suspenso em 'Perfil de risco' (H21) para recalcular a alocação por tipo de FII.")
ws["G32"].font = Font(size=9, italic=True, color="808080")

wb.save(ARQ)
print("Arquivo atualizado: dropdown de perfil + colunas de projeção + gráfico melhorado.")
